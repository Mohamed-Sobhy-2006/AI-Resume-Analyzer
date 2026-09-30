"""Extract plain text from PDF (pypdf) and DOCX (python-docx) resumes."""
import logging
import re
from pathlib import Path

from app.config import settings
from app.utils.file_utils import MSG_CORRUPTED

logger = logging.getLogger(__name__)

MSG_NO_TEXT = ("No readable text was found in this file. "
               "Scanned or image-only resumes are not supported.")


class ResumeParseError(Exception):
    """Raised with a user-friendly message when a resume cannot be read."""


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ResumeParseError("This PDF is password protected. Please upload an unlocked resume.")
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_docx_text(path: Path) -> str:
    from docx import Document
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(stored_name: str, ext: str) -> str:
    """Return cleaned text of an uploaded resume or raise ResumeParseError."""
    path = settings.UPLOAD_DIR / stored_name
    try:
        raw = extract_pdf_text(path) if ext == ".pdf" else extract_docx_text(path)
    except ResumeParseError:
        raise
    except Exception as exc:  # corrupted / unsupported internals of the file
        logger.error("Failed to parse %s: %s", stored_name, exc)
        raise ResumeParseError(MSG_CORRUPTED) from exc
    text = clean_text(raw)
    if len(text) < 30:
        raise ResumeParseError(MSG_NO_TEXT)
    return text
